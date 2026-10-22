"""Requirements-management side: load xlsx, quality checks, traceability, trees, risk, diagrams.
Every *_text function returns plain text (padded markdown tables) so the terminal and the report share one output."""
import re
import openpyxl
from sim import Node, SLOT_FLOW, TICKET_FLOW, LOT_FLOW

ISO25010 = {"Performance": "Performance efficiency", "Availability": "Reliability (availability)",
            "Usability": "Usability", "Security": "Security", "Safety": "Safety", "Scalability": "Performance efficiency (capacity)"}
REFS = {"requirements": ("source", "dependencies"), "use_cases": ("related_requirements",),
        "change_requests": ("requested_by", "affected_requirements"), "hazards": ("related_requirements",),
        "elicitation_sessions": ("participants", "requirements_derived"), "test_cases": ("related_requirements",)}
USE_CASE_TAG = re.compile(r"<<(include|extended by) (UC-\d+)>>")
ID_SHEETS = ("stakeholders", "requirements", "nfrs", "use_cases", "change_requests", "hazards", "test_cases")


def load(path):
    wb = openpyxl.load_workbook(path)
    return {ws.title: [dict(zip([c.value for c in ws[1]], row)) for row in ws.iter_rows(min_row=2, values_only=True) if any(c is not None for c in row)]
            for ws in wb}


def ids(v):
    return [x.strip() for x in str(v).split(";") if x.strip()] if v is not None else []


def table(headers, rows):
    rows = [["" if c is None else str(c).replace("|", "/").replace("\n", " ") for c in r] for r in rows]
    w = [max([len(h)] + [len(r[i]) for r in rows]) for i, h in enumerate(headers)]
    line = lambda r: "| " + " | ".join(c.ljust(w[i]) for i, c in enumerate(r)) + " |"
    return "\n".join([line(headers), "| " + " | ".join("-" * x for x in w) + " |"] + [line(r) for r in rows])


def sil_rule(sev):  # inferred from the sample dataset: severity decides the level
    return {4: "SIL-2", 3: "SIL-1"}.get(sev, "QM")


# ---------------- F1 quality checks ----------------
def quality(data):
    out = []  # (level, id, message)
    known = {r["id"] for s in ID_SHEETS for r in data[s]}
    tested = {q for t in data["test_cases"] for q in ids(t["related_requirements"])}
    with_uc = {q for u in data["use_cases"] for q in ids(u["related_requirements"])}
    for r in data["requirements"]:
        d = r["description"] or ""
        if not r["acceptance_criteria"]:
            out.append(("ERROR", r["id"], "missing acceptance criteria"))
        if d.count(" and ") + d.count(",") >= 3:
            out.append(("WARN", r["id"], "not atomic: description joins several actions (split into separate requirements)"))
        if " shall " not in d:
            out.append(("WARN", r["id"], "description has no 'shall'"))
        if r["id"] not in tested:
            out.append(("WARN", r["id"], "no test case covers this requirement"))
        if r["id"] not in with_uc:
            out.append(("INFO", r["id"], "no use case realizes this requirement"))
        if r["priority"] == "Must" and r["status"] == "Proposed":
            out.append(("INFO", r["id"], "Must requirement still Proposed (needs approval)"))
    for n in data["nfrs"]:
        if not n["metric"] or not n["target_value"]:
            out.append(("ERROR", n["id"], "not measurable: metric or target value missing"))
    for h in data["hazards"]:
        if not h["mitigation"]:
            out.append(("ERROR", h["id"], "no mitigation defined"))
        if not h["related_requirements"]:
            out.append(("WARN", h["id"], "not linked to any requirement"))
        if h["sil"] != sil_rule(h["severity"]):
            out.append(("WARN", h["id"], f"SIL {h['sil']} differs from rule ({sil_rule(h['severity'])} for severity {h['severity']})"))
    actors = {a["name"] for a in data["actors"]}
    for u in data["use_cases"]:
        if u["primary_actor"] not in actors:
            out.append(("ERROR", u["id"], f"unknown primary actor '{u['primary_actor']}'"))
    for sheet, cols in REFS.items():
        for r in data[sheet]:
            for c in cols:
                for x in ids(r[c]):
                    if x not in known:
                        out.append(("ERROR", r["id"], f"{c}: unknown reference {x}"))
    ucs = {u["id"] for u in data["use_cases"]}
    for a in data["test_case_analysis"]:
        if a["test_case_id"] not in known:
            out.append(("ERROR", a["id"], f"unknown test case {a['test_case_id']}"))
        sid = a["source_id"]
        ok = (sid.rsplit("-", 1)[0] in ucs and sid.endswith(("-MAIN", "-ALT"))) if a["source_type"] == "Flow" else sid in known
        if not ok:
            out.append(("ERROR", a["id"], f"unknown source {sid}"))
    flows = {a["source_id"] for a in data["test_case_analysis"] if a["source_type"] == "Flow"}
    for u in data["use_cases"]:
        if f"{u['id']}-MAIN" not in flows:
            out.append(("WARN", u["id"], "main flow has no test case"))
        if ids(u["alternative_flows"]) and f"{u['id']}-ALT" not in flows:
            out.append(("WARN", u["id"], "alternative flows have no test case"))
        for _, other in USE_CASE_TAG.findall(f"{u['main_flow']};{u['alternative_flows']}"):
            if other not in ucs:
                out.append(("ERROR", u["id"], f"include/extend refers to unknown use case {other}"))
    return sorted(out, key=lambda x: ("ERROR", "WARN", "INFO").index(x[0]))


def quality_text(data):
    q = quality(data)
    cnt = {k: sum(1 for x in q if x[0] == k) for k in ("ERROR", "WARN", "INFO")}
    return table(["Level", "ID", "Issue"], q) + f"\n\n{cnt['ERROR']} errors, {cnt['WARN']} warnings, {cnt['INFO']} info."


# ---------------- F2 traceability ----------------
def trace(data):
    def link(rows, col, rid):
        return [x["id"] for x in rows if rid in ids(x[col])]
    rows = []
    for kind, items in (("REQ", data["requirements"]), ("NFR", data["nfrs"])):
        for r in items:
            i = r["id"]
            rows.append([i, link(data["use_cases"], "related_requirements", i), link(data["test_cases"], "related_requirements", i),
                         link(data["hazards"], "related_requirements", i), link(data["change_requests"], "affected_requirements", i),
                         link(data["elicitation_sessions"], "requirements_derived", i)])
    return rows


def trace_text(data):
    rows = trace(data)
    cell = lambda v: ", ".join(v) if v else "-"
    t = table(["Item", "Use case", "Test case", "Hazard", "Change req", "Elicitation"], [[r[0]] + [cell(c) for c in r[1:]] for r in rows])
    reqs = [r for r in rows if r[0].startswith("REQ")]
    pct = lambda col: 100 * sum(1 for r in reqs if r[col]) // len(reqs)
    orphans = [r[0] for r in reqs if not r[2]]
    return t + f"\n\nREQ coverage: use case {pct(1)}%, test case {pct(2)}%, elicitation {pct(5)}%.\nREQ without test: {', '.join(orphans) or 'none'}."


# ---------------- F3 requirement tree + change impact ----------------
def dep_tree(data):
    nodes = {r["id"]: Node(r["id"], "req") for r in data["requirements"]}
    for r in data["requirements"]:
        nodes[r["id"]].title = r["title"]
    roots = []
    for r in data["requirements"]:
        deps = [d for d in ids(r["dependencies"]) if d in nodes]
        for d in deps:
            nodes[d].add(nodes[r["id"]])  # parent = dependency, child = dependent
        if not deps:
            roots.append(nodes[r["id"]])
    return Node("SRS", "root").add(*roots), nodes


def descendants(n):
    out = set()
    for c in n.children:
        out |= {c.name} | descendants(c)
    return out


def tree_text(data):
    root, _ = dep_tree(data)

    def walk(n, prefix=""):
        lines = []
        for i, c in enumerate(n.children):
            last = i == len(n.children) - 1
            lines.append(f"{prefix}{'`-- ' if last else '+-- '}{c.name} {c.title}")
            lines += walk(c, prefix + ("    " if last else "|   "))
        return lines
    return "SRS (child depends on parent)\n" + "\n".join(walk(root))


def impact_text(data):
    _, nodes = dep_tree(data)
    reqs = {r["id"]: r for r in data["requirements"]}
    out = []
    for cr in data["change_requests"]:
        hit = [x for x in ids(cr["affected_requirements"]) if x in nodes]
        imp = set(hit)
        for x in hit:
            imp |= descendants(nodes[x])
        imp = sorted(imp)
        pick = lambda rows, col: sorted(r["id"] for r in rows if set(ids(r[col])) & set(imp))
        out.append(f"{cr['id']} {cr['title']} [{cr['status']}]\n"
                   f"  directly affected : {', '.join(hit)}\n"
                   f"  impacted via deps : {', '.join(imp)}\n"
                   f"  use cases to review: {', '.join(pick(data['use_cases'], 'related_requirements')) or '-'}\n"
                   f"  tests to re-run    : {', '.join(pick(data['test_cases'], 'related_requirements')) or '-'}\n"
                   f"  hazards to review  : {', '.join(pick(data['hazards'], 'related_requirements')) or '-'}\n"
                   f"  complexity of impacted requirements: {sum(reqs[x]['complexity'] for x in imp)}\n"
                   f"  estimate: cost {cr['cost_impact_days']} days, schedule {cr['schedule_impact_days']} days")
    return "\n\n".join(out)


# ---------------- F6 NFR table ----------------
def nfr_text(data, results=None):
    results = results or {}
    return table(["ID", "Title", "Category", "ISO/IEC 25010", "Metric", "Target", "Method", "Simulation result"],
                 [[n["id"], n["title"], n["category"], ISO25010.get(n["category"], "?"), n["metric"] or "(none)", n["target_value"] or "(none)",
                   n["verification_method"], results.get(n["id"], "manual / not simulated")] for n in data["nfrs"]])


# ---------------- F7 risk ----------------
def risk_text(data, demos=None):
    demos = demos or {}
    hz = data["hazards"]
    sevs, liks = (4, 3, 2, 1), (1, 2, 3)
    grid = table(["Severity \\ Likelihood"] + [str(l) for l in liks],
                 [[str(s)] + [", ".join(h["id"] for h in hz if h["severity"] == s and h["likelihood"] == l) for l in liks] for s in sevs])
    ranked = sorted(hz, key=lambda h: -h["severity"] * h["likelihood"])
    rows = [[h["id"], h["hazard"], h["severity"], h["likelihood"], h["severity"] * h["likelihood"], h["sil"],
             h["mitigation"] or "(NONE)", demos.get(h["id"], "-")] for h in ranked]
    return grid + "\n\n" + table(["ID", "Hazard", "Sev", "Lik", "Score", "SIL", "Mitigation", "Mitigation verified in simulation"], rows)


# ---------------- F5 results write-back / test table ----------------
def tests_text(data):
    t = data["test_cases"]
    cnt = {s: sum(1 for x in t if x["status"] == s) for s in ("Passed", "Failed", "Not Run")}
    return (table(["ID", "Title", "Reqs", "Type", "Prio", "Status", "Actual result"],
                  [[x["id"], x["title"], x["related_requirements"], x["type"], x["priority"], x["status"], x["actual_result"]] for x in t])
            + f"\n\nPassed {cnt['Passed']}, Failed {cnt['Failed']}, Not Run {cnt['Not Run']} of {len(t)}.")


def save_results(path, results):  # results: tc id -> (status, actual, by, date)
    wb = openpyxl.load_workbook(path)
    ws = wb["test_cases"]
    col = {c.value: c.column for c in ws[1]}
    for row in ws.iter_rows(min_row=2):
        r = results.get(row[0].value)
        if r:
            for k, v in zip(("status", "actual_result", "executed_by", "executed_date"), r):
                ws.cell(row=row[0].row, column=col[k], value=v)
    wb.save(path)


# ---------------- F8/F9 diagrams (Mermaid) ----------------
def _mid(s):
    return re.sub(r"\W", "", s)


def usecase_mmd(data):
    L = ["flowchart LR", "  subgraph SYS[Smart Parking System]"]
    L += [f'    {_mid(u["id"])}(("{u["title"]}"))' for u in data["use_cases"]]
    L.append("  end")
    for a in data["actors"]:
        L.append(f'  {_mid(a["id"])}(["{a["name"]}"])')
        for u in data["use_cases"]:
            if u["primary_actor"] == a["name"] or a["name"] in f'{u["main_flow"]};{u["alternative_flows"]}':
                L.append(f'  {_mid(a["id"])} --- {_mid(u["id"])}')
    for u in data["use_cases"]:  # <<include UC-x>> in a flow: this UC includes UC-x; <<extended by UC-x>>: UC-x extends this UC
        for kind, other in USE_CASE_TAG.findall(f'{u["main_flow"]};{u["alternative_flows"]}'):
            L.append(f'  {_mid(u["id"])} -.->|include| {_mid(other)}' if kind == "include" else f'  {_mid(other)} -.->|extend| {_mid(u["id"])}')
    return "\n".join(L)


def state_mmd():
    def sm(flow, start, end=None):
        L = ["stateDiagram-v2", f"  [*] --> {start}"] + [f"  {a} --> {b}" for a, bs in flow.items() for b in sorted(bs)]
        return "\n".join(L + ([f"  {end} --> [*]"] if end else []))
    return sm(SLOT_FLOW, "Free"), sm(TICKET_FLOW, "Active", "Closed"), sm(LOT_FLOW, "Normal")


DFD_MMD = """flowchart LR
  Driver([Driver]) -->|frame, type| P7(7 Recognize plate)
  Camera([Entry Camera]) -->|plate image| P7
  P7 -->|plate| P1(1 Entry and nearest slot)
  P1 -->|ticket, route| Driver
  Driver -->|ticket| P2(2 Exit and dynamic fee)
  P2 <-->|payment| PG([Payment Gateway])
  Driver -->|reservation request| P3(3 Reserve)
  Driver -->|plate| P4(4 Locate)
  Attendant([Gate Attendant]) -->|plate| P5(5 Lost ticket)
  Admin([Administrator]) -->|lock / unlock| P6(6 Zone control)
  Sensor([Slot Sensor]) -->|readings| P8(8 Monitor slots)
  P8 -->|alerts| Attendant
  Alarm([Fire Alarm]) -->|signal| P9(9 Evacuation)
  Admin -->|trigger| P9
  P1 <--> D1[(D1 Slots tree)]
  P2 <--> D1
  P3 <--> D1
  P4 --> D1
  P6 --> D1
  P8 <--> D1
  P9 --> D1
  P1 --> D2[(D2 Tickets)]
  P2 <--> D2
  P5 <--> D2
  P8 --> D2
  P1 --> D3[(D3 Audit log)]
  P2 --> D3
  P5 --> D3
  P6 --> D3
  P8 --> D3
  P9 --> D3
  Admin -->|account, password| P10(10 Log in)
  Manager([Parking Manager]) -->|account, password| P10
  P10 --> D4[(D4 Accounts)]
  Admin -->|rate, new zone| P11(11 Configure)
  P11 --> D1
  P11 --> D3
  Manager -->|period| P12(12 Report)
  P12 -->|revenue, audit| Manager
  P12 --> D3
  Driver -->|plate| P13(13 Cancel reservation)
  P13 --> D1
  Driver -->|look| P14(14 Availability)
  P14 --> D1
  Camera2([Exit Camera]) -->|exit plate| P15(15 Verify plate at exit)
  P15 --> P2
  P15 --> D3
  Attendant -->|approval, override request| P16(16 Approve and override)
  Manager -->|decision| P16
  P16 --> D2
  P16 --> D3
  Admin -->|pass, VIP| P17(17 Passes)
  P17 --> D5[(D5 Passes)]
  P1 --> D5
  Charger([Charging Station]) -->|kWh| P18(18 EV charging)
  P18 --> D2
  Server([Central Server]) <-->|sync, reservations, passes| P19(19 Offline mode and sync)
  P19 --> D3"""


def lot_mmd(root):
    L, n = ["flowchart TD"], [0]

    def walk(node):
        n[0] += 1
        me = f"n{n[0]}"
        L.append(f'  {me}["{node.name}"]')
        for c in node.children:
            L.append(f"  {me} --> {walk(c)}")
        return me
    walk(root)
    return "\n".join(L)


def dep_mmd(data):
    return "\n".join(["flowchart LR"] + [f'  {_mid(d)} --> {_mid(r["id"])}' for r in data["requirements"] for d in ids(r["dependencies"])])


def grid_text(data):
    lv = ("High", "Medium", "Low")
    return table(["Influence \\ Interest"] + list(lv),
                 [[i] + [", ".join(s["id"] for s in data["stakeholders"] if s["influence"] == i and s["interest"] == j) for j in lv] for i in lv])


def dataset_text(data):
    return "\n\n".join(f"### {name}\n\n" + table(list(rows[0]), [list(r.values()) for r in rows]) for name, rows in data.items() if rows)
